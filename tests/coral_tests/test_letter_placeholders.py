"""Letter placeholder parsing, without Django: the two pure helpers are lifted out of
file_template.py. `python tests/coral_tests/test_letter_placeholders.py`
"""

import ast
import os
import re

from docx import Document

SOURCE = os.path.join(os.path.dirname(__file__), '..', '..', 'coral', 'views', 'file_template.py')
tree = ast.parse(open(SOURCE).read())
wanted = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in ('placeholders', 'split_placeholders')]
assert len(wanted) == 2
namespace = {'re': re}
exec(compile(ast.Module(body=wanted, type_ignores=[]), SOURCE, 'exec'), namespace)
placeholders, split_placeholders = namespace['placeholders'], namespace['split_placeholders']

doc = Document()
doc.add_paragraph('Dear <owner>, re <monument_name>')
doc.add_table(rows=1, cols=1).cell(0, 0).text = '<owner__street_value>'
doc.sections[0].header.paragraphs[0].text = '<today>'
doc.sections[0].footer.paragraphs[0].text = '<user>'
assert placeholders(doc) == {'owner', 'monument_name', 'owner__street_value', 'today', 'user'}, placeholders(doc)

own, related, specials = split_placeholders(
    {'owner', 'owner__street_value', 'owner__county_value', 'street_value', 'today'},
    ['owner'], {'today': 'today'})
assert own == {'owner', 'street_value'}, own
assert related == {'owner': {'street_value', 'county_value'}}, related
assert specials == {'today'}, specials

print('ok')
